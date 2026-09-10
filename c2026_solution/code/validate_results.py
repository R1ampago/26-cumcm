"""Structural and numerical checks for generated result workbooks."""
from __future__ import annotations

import json
import sys
from datetime import date
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

HERE = Path(__file__).resolve()
ROOT = HERE.parents[2]
if str(HERE.parent) not in sys.path:
    sys.path.insert(0, str(HERE.parent))

from data_io import DT_HOURS, interval_labels  # noqa: E402


REPORT_DATES = [date.fromisoformat("2025-02-01")]
# Avoid depending on a date utility in the validator's public output.
from datetime import timedelta  # noqa: E402
REPORT_DATES = [REPORT_DATES[0] + timedelta(days=i) for i in range(334)]


def _as_float(v, name: str) -> float:
    if v is None or isinstance(v, str):
        raise AssertionError(f"{name} is not numeric: {v!r}")
    value = float(v)
    if not np.isfinite(value):
        raise AssertionError(f"{name} is not finite")
    return value


def validate_plan_sheet(ws, expected_dates: list[date]) -> None:
    if ws.max_column < 145:
        raise AssertionError(f"{ws.title}: expected 144 interval columns")
    labels = interval_labels()
    got_labels = [ws.cell(1, j).value for j in range(2, 146)]
    if got_labels != labels:
        raise AssertionError(f"{ws.title}: interval labels do not use the audited endpoint convention")
    got_dates = [ws.cell(i, 1).value.date() for i in range(2, len(expected_dates) + 2)]
    if got_dates != expected_dates:
        raise AssertionError(f"{ws.title}: date rows are not exactly Feb 1--Dec 31")
    for i in range(2, len(expected_dates) + 2):
        for j in range(2, 146):
            if _as_float(ws.cell(i, j).value, f"{ws.title}!{i},{j}") < -1e-8:
                raise AssertionError(f"{ws.title}: negative purchase")


def validate_storage_sheet(ws, expected_days: int) -> None:
    if ws.max_row != 1 + expected_days * 6:
        raise AssertionError(f"{ws.title}: expected {1 + expected_days * 6} rows, got {ws.max_row}")
    for d in range(expected_days):
        start = 2 + d * 6
        if ws.cell(start, 1).value is None:
            raise AssertionError(f"{ws.title}: missing date at day {d}")
        for r in range(start, start + 6):
            _as_float(ws.cell(r, 3).value, f"charge row {r}")
            _as_float(ws.cell(r, 4).value, f"discharge row {r}")
            if float(ws.cell(r, 3).value) < -1e-8 or float(ws.cell(r, 4).value) < -1e-8:
                raise AssertionError(f"{ws.title}: negative storage amount")
        for r in range(start, start + 6):
            if ws.cell(r, 5).value in {"0:00", "24:00"}:
                soc = _as_float(ws.cell(r, 6).value, f"SOC row {r}")
                if not 1200 - 1e-5 <= soc <= 10800 + 1e-5:
                    raise AssertionError(f"{ws.title}: SOC outside bounds")


def validate_emergency_sheet(ws) -> int:
    count = 0
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 2).value in (None, "⁝"):
            continue
        energy = _as_float(ws.cell(r, 3).value, f"emergency row {r}")
        if energy < -1e-8:
            raise AssertionError("negative emergency purchase")
        count += 1
    return count


def main() -> None:
    result_dir = ROOT / "c2026_solution" / "results"
    expected_dates = REPORT_DATES
    checks: dict[str, object] = {}

    wb = load_workbook(result_dir / "result1.xlsx", data_only=True)
    ws = wb["计划购电量"]
    if ws.max_row != 145 or ws.max_column < 2:
        raise AssertionError("result1 plan shape mismatch")
    if [ws.cell(i, 1).value for i in range(2, 146)] != interval_labels():
        raise AssertionError("result1 labels mismatch")
    for i in range(2, 146):
        if _as_float(ws.cell(i, 2).value, f"result1 row {i}") < -1e-8:
            raise AssertionError("result1 negative purchase")
    checks["result1"] = "ok"

    for name, expected_sheets in [
        ("result2.xlsx", ["计划购电量", "充放电量", "紧急购电量"]),
        ("result4-2.xlsx", ["计划购电量", "充放电量", "紧急购电量"]),
        ("result3.xlsx", ["计划购电量", "调整购电量", "充放电量", "紧急购电量"]),
        ("result4-3.xlsx", ["计划购电量", "调整购电量", "充放电量", "紧急购电量"]),
    ]:
        wb = load_workbook(result_dir / name, data_only=True)
        if wb.sheetnames != expected_sheets:
            raise AssertionError(f"{name}: sheet names mismatch")
        validate_plan_sheet(wb["计划购电量"], expected_dates)
        if "调整购电量" in wb.sheetnames:
            validate_plan_sheet(wb["调整购电量"], expected_dates)
        validate_storage_sheet(wb["充放电量"], len(expected_dates))
        event_count = validate_emergency_sheet(wb["紧急购电量"])
        checks[name] = {"status": "ok", "emergency_segments": event_count}

    out = ROOT / "c2026_solution" / "audit" / "result_validation.json"
    out.write_text(json.dumps(checks, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(checks, ensure_ascii=False, indent=2))
    print(f"wrote {out}")


if __name__ == "__main__":
    main()
