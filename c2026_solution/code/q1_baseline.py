"""Question 1: representative-day linear-programming baseline.

Run from the repository root:
    python c2026_solution/code/q1_baseline.py
"""
from __future__ import annotations

import json
import shutil
import sys
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

HERE = Path(__file__).resolve()
ROOT = HERE.parents[2]
CODE = HERE.parent
if str(CODE) not in sys.path:
    sys.path.insert(0, str(CODE))

from data_io import DT_HOURS, interval_labels, load_attachments, typical_arrays  # noqa: E402
from dispatch_lp import StorageParams, solve_dispatch  # noqa: E402


def fill_result_workbook(dispatch, output_path: Path) -> None:
    source = ROOT / "CUMCM2026Problems" / "C题" / "附件" / "附件5" / "result1.xlsx"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, output_path)
    wb = load_workbook(output_path)
    plan = wb["计划购电量"]
    # The supplied template labels the first column one interval late
    # (0:10-0:20 ... 0:00+1-0:10+1), while the question's 0:00/24:00
    # boundary conditions require the endpoint convention 0:00-0:10 ...
    # 23:50-24:00.  We correct labels in generated workbooks rather than
    # silently shifting the physical solution.
    for k, label in enumerate(interval_labels(), start=2):
        plan.cell(row=k, column=1).value = label
        plan.cell(row=k, column=2).value = round(float(dispatch.grid_energy[k - 2]), 6)

    storage = wb["充放电量"]
    for group in range(6):
        a = group * 24
        b = (group + 1) * 24
        storage.cell(row=group + 2, column=2).value = round(float(dispatch.charge_energy[a:b].sum()), 6)
        storage.cell(row=group + 2, column=3).value = round(float(dispatch.discharge_energy[a:b].sum()), 6)
    storage.cell(row=2, column=5).value = round(float(dispatch.soc[0]), 6)
    storage.cell(row=3, column=5).value = round(float(dispatch.soc[-1]), 6)
    wb.save(output_path)


def main() -> None:
    data = load_attachments(ROOT)
    price, load, pv = typical_arrays(data)
    params = StorageParams()
    dispatch = solve_dispatch(
        load,
        pv,
        price,
        soc_initial=params.initial_soc,
        soc_terminal=params.initial_soc,
        params=params,
    )
    out_dir = ROOT / "c2026_solution" / "results"
    out_dir.mkdir(parents=True, exist_ok=True)
    output_path = out_dir / "result1.xlsx"
    fill_result_workbook(dispatch, output_path)

    summary = {
        "objective_yuan": float(dispatch.objective),
        "grid_energy_kwh": float(dispatch.grid_energy.sum()),
        "charge_energy_kwh": float(dispatch.charge_energy.sum()),
        "discharge_energy_kwh": float(dispatch.discharge_energy.sum()),
        "curtail_energy_kwh": float(dispatch.energy(dispatch.curtail_power).sum()),
        "soc_initial_kwh": float(dispatch.soc[0]),
        "soc_terminal_kwh": float(dispatch.soc[-1]),
        "soc_min_kwh": float(dispatch.soc.min()),
        "soc_max_kwh": float(dispatch.soc.max()),
        "max_charge_power_kw": float(dispatch.charge_power.max()),
        "max_discharge_power_kw": float(dispatch.discharge_power.max()),
        "simultaneous_charge_discharge_intervals": int(
            np.sum((dispatch.charge_power > 1e-7) & (dispatch.discharge_power > 1e-7))
        ),
        "dt_hours": DT_HOURS,
        "efficiency": params.eta,
        "dispatch_cost_definition": "sum(price_k * grid_power_k * dt)",
    }
    (out_dir / "q1_summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    print(f"wrote {output_path}")


if __name__ == "__main__":
    main()
